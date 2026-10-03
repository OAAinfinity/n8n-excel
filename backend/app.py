from fastapi import Body, FastAPI, HTTPException
from fastapi.responses import FileResponse

from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path

app = FastAPI()


@app.post("/generate-excel")
def generate_excel(data: list | dict | None = Body(default=None)):
    if isinstance(data, dict):
        if "data" not in data:
            raise HTTPException(status_code=400, detail="Request body must be a list or an object with a 'data' array.")
        data = data["data"]

    if not isinstance(data, list):
        raise HTTPException(status_code=400, detail="Request body must be a list or an object with a 'data' array.")

    base_dir = Path(__file__).resolve().parent
    template_path = base_dir / "templates" / "FAAIRMOUNT_TEMPLATE.xlsx"
    output_dir = base_dir / "generated"
    output_dir.mkdir(parents=True, exist_ok=True)

    if not template_path.exists():
        raise HTTPException(status_code=500, detail=f"Template file not found: {template_path}")

    wb = load_workbook(template_path)
    ws = wb.active

    row = 2

    for item in data:
        if not isinstance(item, dict):
            raise HTTPException(status_code=400, detail="Each entry in data must be an object.")

        ws[f"A{row}"] = item.get("date", "")
        ws[f"B{row}"] = item.get("type", "")
        ws[f"C{row}"] = item.get("particular", "")
        ws[f"E{row}"] = item.get("quantity", 0)
        ws[f"F{row}"] = item.get("rate", 0)

        ws[f"G{row}"] = f"=E{row}*F{row}"

        row += 1

    filename = output_dir / f"Report_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"

    wb.save(filename)

    return FileResponse(
        str(filename),
        filename=filename.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )